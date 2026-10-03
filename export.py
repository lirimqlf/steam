"""export.py — génération de l'Excel (colonnes de points = critères activés seulement)."""
import io
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import scoring


def _d(ts):
    return datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d") if ts else ""


def build_xlsx(rows: list, cfg: dict, statuses: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    on = [k for k in cfg if cfg[k]["enabled"]]
    head = ["Jeu", "Statut", "Score", "Joueurs", "Reviews", "Dernière review", "Dernière news",
            "Inactif (jours)", "Sortie", "Développeur", "Éditeur", "Email", "Site / support",
            "Prix ($)", "Steam", "Discord"] + [f"pts · {scoring.LABELS[k]}" for k in on]
    ws.append(head)
    for r in rows:
        c = r.get("contact") or ""
        email = r.get("email") or (c if "@" in c else "")
        site = r.get("website") or (c if c and "@" not in c else "")
        price = "Free" if r.get("is_free") else (r["price"] / 100 if r.get("price") is not None else "")
        bd = scoring.breakdown(r, cfg)
        ws.append([r["name"], r["status"], r.get("score") or 0, r.get("current_players"),
                   r.get("total_reviews"), _d(r.get("last_review_ts")), _d(r.get("last_news_ts")),
                   scoring.inactive_days(r), r.get("release_date") or "", r.get("developers") or "",
                   r.get("publishers") or "", email, site, price, "Steam", "Claim"] + [bd[k] for k in on])
        i = ws.max_row
        ws.cell(i, 15).hyperlink = f"https://store.steampowered.com/app/{r['steam_appid']}"
        ws.cell(i, 16).hyperlink = f"https://discord.com/developers/applications/{r['discord_id']}"
        for col in (15, 16):
            ws.cell(i, col).font = Font(color="1F5FFF", underline="single")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30 if i in (1, 10, 11, 12, 13) else 14
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    cs = wb.create_sheet("Critères")
    cs.append(["Critère", "Activé", "Points"])
    for k, v in cfg.items():
        cs.append([v["label"], "oui" if v["enabled"] else "non", v["points"]])
    cs.append([])
    cs.append(["Statuts exportés", ", ".join(statuses)])
    cs.append(["Score", "normalisé sur 100 selon le total des critères activés"])
    cs.append(["Généré le", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    for c in cs[1]:
        c.font = Font(bold=True)
    cs.column_dimensions["A"].width = 52
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
